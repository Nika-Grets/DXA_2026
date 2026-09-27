from pathlib import Path
import struct, json, hashlib
import numpy as np
from openpyxl import load_workbook
from scipy import ndimage
import sys
ROOT=Path('/mnt/data/dxa_work')
DROOT=ROOT/'.training_dicom'

def read_dcm(path):
    b=Path(path).read_bytes()
    if b[128:132] != b'DICM': raise ValueError('no DICM')
    # file meta group length (0002,0000), explicit VR
    g,e=struct.unpack_from('<HH',b,132); vr=b[136:138]
    if (g,e)!=(2,0): raise ValueError('no meta group length')
    gl=struct.unpack_from('<I',b,140)[0]
    off=144+gl
    vals={}
    while off+8<=len(b):
        g,e,l=struct.unpack_from('<HHI',b,off); off+=8
        if l==0xffffffff: raise ValueError('undefined length')
        v=b[off:off+l]
        if (g,e) in [(0x0020,0x000d),(0x0008,0x0018),(0x0028,0x0010),(0x0028,0x0011),(0x0028,0x0100),(0x0028,0x0101),(0x0028,0x0102),(0x0028,0x0103),(0x0028,0x0002),(0x0028,0x0004),(0x7fe0,0x0010)]:
            vals[(g,e)] = v
        off += l
        if (g,e)==(0x7fe0,0x0010): break
    rows=struct.unpack('<H',vals[(0x0028,0x0010)][:2])[0]
    cols=struct.unpack('<H',vals[(0x0028,0x0011)][:2])[0]
    bits=struct.unpack('<H',vals[(0x0028,0x0100)][:2])[0]
    rep=struct.unpack('<H',vals[(0x0028,0x0103)][:2])[0]
    pix=vals[(0x7fe0,0x0010)]
    if bits==8: arr=np.frombuffer(pix,dtype=np.uint8,count=rows*cols).reshape(rows,cols)
    elif bits==16:
        arr=np.frombuffer(pix,dtype='<i2' if rep else '<u2',count=rows*cols).reshape(rows,cols)
    else: raise ValueError(f'bits {bits}')
    study=vals.get((0x0020,0x000d),b'').decode('ascii','ignore').rstrip('\x00 ')
    sop=vals.get((0x0008,0x0018),b'').decode('ascii','ignore').rstrip('\x00 ')
    return study,sop,arr

# labels
xlsx=next(DROOT.rglob('*.xlsx'))
ws=load_workbook(xlsx,data_only=True).active
labels={}
for r in range(3,ws.max_row+1):
    study=ws.cell(r,2).value
    if not study: continue
    labels[str(study)]={'source_folder_id':str(study),'spine_position':ws.cell(r,3).value,'spine_axis':ws.cell(r,4).value,'spine_artifact':ws.cell(r,5).value}

# map workbook source_folder_id -> canonical DICOM StudyInstanceUID
manifest=json.loads((ROOT/'artifacts/canonical_manifest.json').read_text())
source_to_uid={str(x['source_folder_id']):str(x['study_uid']) for x in manifest['studies']}
labels={source_to_uid[k]:v for k,v in labels.items() if k in source_to_uid}

# load images grouped by study
studies={}
for p in DROOT.rglob('*.dcm'):
    try:
        study,sop,arr=read_dcm(p)
    except Exception as e:
        print('ERR',p,e); continue
    studies.setdefault(study,[]).append((sop,p,arr))
print('studies',len(studies),'images',sum(map(len,studies.values())))

# try production router
sys.path.insert(0,str(ROOT/'backend'))
from app.ml.dexa_baseline import classify_study, _benchmark_features, _criterion_prob, _normalize
rows=[]
for study, items in studies.items():
    # stable order
    items=sorted(items,key=lambda x:x[0])
    ims=[x[2] for x in items]
    routes=classify_study(ims)
    for it,rt in zip(items,routes):
        if rt.region=='spine':
            x=_normalize(it[2]);
            feat=_benchmark_features(it[2])
            # model direct, avoid cache settings
            import joblib
            model=joblib.load(ROOT/'artifacts/models/spine_artifact.joblib')
            p=float(model.predict_proba(feat)[0,1])
            # handcrafted stats
            xx=x
            q={
              'mean':float(xx.mean()), 'std':float(xx.std()), 'bright97':float((xx>.97).mean()),
              'bright99':float((xx>.99).mean()), 'dark01':float((xx<.01).mean()),
              'grad':float(np.abs(np.diff(xx,axis=0)).mean()+np.abs(np.diff(xx,axis=1)).mean()),
              'sym':float(np.corrcoef(xx.ravel(),np.fliplr(xx).ravel())[0,1]),
              'sat_components':0, 'sat_max_area':0, 'sat_area':0,
            }
            mask=xx>.97
            lab,n=ndimage.label(mask)
            sizes=np.bincount(lab.ravel())[1:] if n else np.array([])
            if sizes.size:
                q['sat_components']=int((sizes>=4).sum()); q['sat_max_area']=int(sizes.max()); q['sat_area']=int(sizes[sizes>=4].sum())
            rows.append({'study':study,'sop':it[0],'path':str(it[1]),'label':labels.get(study,{}).get('spine_artifact'), 'p_art':p, **q})

out=Path(ROOT/'artifacts/spine_artifact_diagnostic.json')
out.write_text(json.dumps(rows,ensure_ascii=False,indent=2))
print('spine images',len(rows),'positive studies',sum(1 for s,v in labels.items() if v.get('spine_artifact')==1))
# summarize by label, image-level
for lab in [0,1]:
    a=[r for r in rows if r['label']==lab]
    print('\nLABEL',lab,'images',len(a),'studies',len(set(r['study'] for r in a)))
    for k in ['p_art','mean','std','bright97','bright99','dark01','grad','sym','sat_components','sat_max_area','sat_area']:
        z=np.array([r[k] for r in a],float)
        print(k,'median',np.median(z),'p90',np.quantile(z,.9),'max',np.max(z))
# study aggregate max p and handcrafted max
print('\nSTUDY AGGREGATE')
for lab in [0,1]:
    a=[]
    for s in labels:
        if labels[s].get('spine_artifact')==lab and s in studies:
            rr=[r for r in rows if r['study']==s]
            if rr:a.append(max(rr,key=lambda r:r['p_art']))
    print('label',lab,'n',len(a),'max p median/p90/max',np.median([r['p_art'] for r in a]),np.quantile([r['p_art'] for r in a],.9),max(r['p_art'] for r in a))
# top 20 p
print('\nTOP P')
for r in sorted(rows,key=lambda r:r['p_art'],reverse=True)[:20]: print(r['label'],f"{r['p_art']:.4f}",r['study'],r['sop'],r['sat_components'],r['sat_max_area'],r['bright97'])
