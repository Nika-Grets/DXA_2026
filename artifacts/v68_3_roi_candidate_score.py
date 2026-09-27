from __future__ import annotations
import json, sys, warnings
from pathlib import Path
import numpy as np
from sklearn.metrics import roc_auc_score, average_precision_score
from sklearn.model_selection import StratifiedGroupKFold
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
TRAIN=Path('/mnt/data/train_v68')
sys.path.insert(0, str(ROOT/'backend'))
sys.path.insert(0, '/mnt/data')
from app.ml.dexa_baseline import classify_study, _normalize
from app.ml.hip_localizer import localize_hip
from mini_dicom import read_pixels

MAN=json.loads((ROOT/'artifacts/canonical_manifest.json').read_text())
BY={s['source_folder_id']:s for s in MAN['studies']}
STUDY_DIR=next(TRAIN.glob('#U0418*'))
XLSX=next(TRAIN.glob('*.xlsx'))

# q60_sx = horizontal spread (std x) of pixels above the 60th percentile
# in the anatomy-localizer's shaft-anchored proximal-femur crop.
# It is a candidate score only; no production threshold is set here.
def feature_q60_sx(image, side):
    x=_normalize(image); h,w=x.shape
    loc=localize_hip(x,side)
    score,cx,sty,sbw,sbh=loc.shaft_score, loc.bone_bbox[0]*w, 0, 0, 0
    # Reuse the same shaft candidate exposed by the experimental V6.8 anatomy code.
    from app.ml.hip_localizer import _candidate_mask, _shaft_candidate
    mask=_candidate_mask(x); shaft=_shaft_candidate(mask,x,side)
    if shaft is None: return None
    _,cx,sty,_,_=shaft
    xx=(x if side=='right' else np.fliplr(x)).astype(np.float32)
    if side=='left': cx=w-cx
    x0=max(0,int(cx-0.28*w)); x1=min(w,int(cx+0.28*w))
    y0=max(0,int(sty-0.34*h)); y1=min(h,int(sty+0.22*h))
    crop=xx[y0:y1,x0:x1]
    if crop.size<100: return None
    crop=__import__('cv2').resize(crop,(40,48),interpolation=__import__('cv2').INTER_AREA)
    m=crop>=np.percentile(crop,60)
    yy,xx2=np.nonzero(m)
    return float(xx2.std()/40) if len(xx2)>5 else None

wb=load_workbook(XLSX,data_only=True); ws=wb.active
records=[]
for r in ws.iter_rows(min_row=3, values_only=True):
    src=r[1]
    if src not in BY: continue
    if r[6] is None and r[8] is None: continue
    s=BY[src]
    paths=[STUDY_DIR/im['relative_path'] for ser in s['series'] for im in ser['images']]
    images=[read_pixels(str(p))[0].astype(np.float32) for p in paths]
    preds=classify_study(images)
    for side,pc,rc in [('right',r[5],r[6]),('left',r[7],r[8])]:
        if pc is None or rc is None: continue
        pp=[p for p in preds if p.region==f'femur_{side}']
        if not pp: continue
        v=feature_q60_sx(images[pp[0].image_index],side)
        if v is not None:
            records.append({'study':s['study_uid'],'side':side,'roi':int(rc),'q60_sx':v})

y=np.array([r['roi'] for r in records]); p=np.array([r['q60_sx'] for r in records]); g=np.array([r['study'] for r in records])
result={'n':len(y),'positive':int(y.sum()),'roc_auc':float(roc_auc_score(y,p)),'pr_auc':float(average_precision_score(y,p))}
# Study-level bootstrap CI.
rng=np.random.default_rng(42); studies=np.unique(g); boots=[]
for _ in range(3000):
    sampled=rng.choice(studies,size=len(studies),replace=True)
    idx=np.concatenate([np.where(g==s)[0] for s in sampled])
    if len(np.unique(y[idx]))<2: continue
    boots.append(roc_auc_score(y[idx],p[idx]))
result['bootstrap_95ci']=[float(x) for x in np.percentile(boots,[2.5,97.5])]
result['by_side']={}
for side in ('right','left'):
    m=np.array([r['side']==side for r in records])
    result['by_side'][side]={'n':int(m.sum()),'positive':int(y[m].sum()),'roc_auc':float(roc_auc_score(y[m],p[m]))}
# Fixed exploratory thresholds; explicitly not selected for production.
result['threshold_diagnostics']={}
for t in (0.23,0.235,0.24,0.25,0.27):
    pred=p>=t
    tp=int(((pred)&(y==1)).sum()); fp=int(((pred)&(y==0)).sum()); fn=int(((~pred)&(y==1)).sum()); tn=int(((~pred)&(y==0)).sum())
    result['threshold_diagnostics'][str(t)]={'tp':tp,'fp':fp,'fn':fn,'tn':tn,'sensitivity':tp/(tp+fn),'specificity':tn/(tn+fp)}

out=ROOT/'artifacts'/'v68_3_roi'; out.mkdir(exist_ok=True)
(out/'results.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
(out/'records.json').write_text(json.dumps(records,ensure_ascii=False,indent=2))
print(json.dumps(result,ensure_ascii=False,indent=2))
