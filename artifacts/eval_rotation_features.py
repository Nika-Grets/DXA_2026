import sys,json,cv2
from pathlib import Path
import numpy as np
from sklearn.metrics import roc_auc_score
from openpyxl import load_workbook
sys.path.insert(0,'/mnt/data/v63/backend')
from app.ml.dexa_baseline import classify_study,_normalize,_crop_black_edge_mask
from app.ml.hip_localizer import localize_hip
ROOT=[p for p in Path('/mnt/data/dexa_raw').iterdir() if p.is_dir()][0]; CACHE=Path('/mnt/data/work_spine/dcm_cache');mp=json.load(open('/mnt/data/work_spine/dcm_path_hash.json'));split=json.load(open('/mnt/data/v63/artifacts/study_split.json'))['split']; ws=load_workbook(next(Path('/mnt/data/dexa_raw').glob('*.xlsx')),data_only=True).active

def feats(a,side,loc):
 x=_normalize(a); h,w=x.shape; bx,by,bw,bh=loc.bone_bbox
 x0=max(0,int((bx-.10*bw)*w)); x1=min(w,int((bx+1.10*bw)*w)); y0=max(0,int((by-.10*bh)*h)); y1=min(h,int((by+1.10*bh)*h)); crop=x[y0:y1,x0:x1]
 m=(crop>np.percentile(crop,75)).astype(np.uint8); m=cv2.morphologyEx(m,cv2.MORPH_CLOSE,np.ones((5,5),np.uint8))
 H,W=m.shape
 widths=[]; med=[]; lat=[]
 for yy in range(int(.25*H),int(.85*H)):
  xs=np.where(m[yy]>0)[0]
  if xs.size:
   lo,hi=xs.min(),xs.max(); widths.append((hi-lo+1)/W)
   if side=='right': med.append((W-lo)/W); lat.append(hi/W)
   else: med.append((hi+1)/W); lat.append((W-lo)/W)
 if not widths:return [0]*7
 widths=np.array(widths);med=np.array(med);lat=np.array(lat)
 return [float(np.max(widths)),float(np.mean(widths)),float(np.percentile(widths,90)),float(np.max(med)),float(np.mean(med)),float(np.percentile(med,90)),float(np.max(lat))]
rows=[]
for r in ws.iter_rows(min_row=3,values_only=True):
 if not r[1] or split.get(str(r[1]))!='validation':continue
 ims=[];seen=set()
 for fp in sorted((ROOT/str(r[1])).rglob('*.dcm')):
  k=mp.get(str(fp));
  if k and k not in seen:seen.add(k);ims.append(np.load(CACHE/f'{k}.npy').astype(np.float32))
 for side,col in [('right',5),('left',7)]:
  if r[col] is None:continue
  for pr in classify_study(ims):
   if pr.region==f'femur_{side}':
    loc=localize_hip(ims[pr.image_index],side); fs=feats(ims[pr.image_index],side,loc)
    rows.append((int(r[col]),fs));break
Y=np.array([r[0] for r in rows]); X=np.array([r[1] for r in rows])
print('n',len(Y),'pos',Y.sum())
for i,n in enumerate(['maxw','meanw','p90w','maxmed','meanmed','p90med','maxlat']):
 s=X[:,i]; print(n,round(max(roc_auc_score(Y,s),roc_auc_score(Y,-s)),3), 'raw',round(roc_auc_score(Y,s),3))
